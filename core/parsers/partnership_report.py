# -*- coding: utf-8 -*-
"""제휴(인플루언서 방문형 체험단) 오픈 보고서 파서.

'인스타 집행 보고서'(인플루언서 본인 계정 게시물 성과)와
'블로그 집행 보고서'(네이버 블로그 성과) 두 시트에서 개별 행을 읽는다.
'합계' 행은 우리가 새로 합산하므로 원본의 합계 행은 건너뛴다.
이 데이터는 청년상인 자체 채널 콘텐츠가 아니라 외부 인플루언서 계정의
게시물이므로 발행리스트와 매칭하지 않고 별도 '제휴' 시트로 취합한다.
"""
from dataclasses import dataclass
import datetime as dt
import openpyxl


@dataclass
class InstaPartnerRow:
    open_date: dt.date
    handle: str
    summary: str
    followers: int
    views: int
    likes: int
    comments: int
    note: str
    url: str = None


@dataclass
class BlogPartnerRow:
    open_date: dt.date
    blogger: str
    title: str
    daily_visitors: int
    open_day_visitors: int
    pv: int
    comments: int
    likes: int
    note: str
    url: str = None


def _to_date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    return None


def _find_header_row(ws, must_contain: str, search_rows=20):
    for r in range(1, min(ws.max_row, search_rows) + 1):
        values = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
        if must_contain in values:
            return r, {v: c for c, v in enumerate(values, start=1) if v}
    return None, {}


def _get_url(ws, row, col):
    if not col:
        return None
    cell = ws.cell(row=row, column=col)
    return cell.hyperlink.target if cell.hyperlink else None


def parse_partnership_report(file_like) -> dict:
    wb = openpyxl.load_workbook(file_like, data_only=True)
    insta_rows, blog_rows = [], []

    if "인스타 집행 보고서" in wb.sheetnames:
        ws = wb["인스타 집행 보고서"]
        header_row, header = _find_header_row(ws, "오픈 일")
        if header_row:
            def get(row, name):
                c = header.get(name)
                return ws.cell(row=row, column=c).value if c else None
            last_date = None
            for r in range(header_row + 1, ws.max_row + 1):
                handle = get(r, "인스타 아이디")
                if not handle or handle == "합 계":
                    continue
                # 오픈 일은 병합 셀이라 배치의 두 번째 행부터 값이 비어 있음 -> 이전 값 유지
                last_date = _to_date(get(r, "오픈 일")) or last_date
                insta_rows.append(InstaPartnerRow(
                    open_date=last_date,
                    handle=str(handle).strip(),
                    summary=str(get(r, "내용 요약") or "").strip(),
                    followers=get(r, "팔로워 수"),
                    views=get(r, "조회 수"),
                    likes=get(r, "좋아요 수"),
                    comments=get(r, "댓글 수"),
                    note=str(get(r, "비고") or "").strip(),
                    url=_get_url(ws, r, header.get("URL")),
                ))

    if "블로그 집행 보고서" in wb.sheetnames:
        ws = wb["블로그 집행 보고서"]
        header_row, header = _find_header_row(ws, "오픈 일")
        if header_row:
            def get(row, name):
                c = header.get(name)
                return ws.cell(row=row, column=c).value if c else None
            last_date = None
            for r in range(header_row + 1, ws.max_row + 1):
                blogger = get(r, "진행 블로거")
                if not blogger or blogger == "합 계":
                    continue
                last_date = _to_date(get(r, "오픈 일")) or last_date
                blog_rows.append(BlogPartnerRow(
                    open_date=last_date,
                    blogger=str(blogger).strip(),
                    title=str(get(r, "제목") or "").strip(),
                    daily_visitors=get(r, "일평균\n방문자수"),
                    open_day_visitors=get(r, "오픈일\n방문자수"),
                    pv=get(r, "PV수치"),
                    comments=get(r, "댓글 수"),
                    likes=get(r, "공감 수"),
                    note=str(get(r, "비고") or "").strip(),
                    url=_get_url(ws, r, header.get("URL")),
                ))

    return {"인스타": insta_rows, "블로그": blog_rows}
