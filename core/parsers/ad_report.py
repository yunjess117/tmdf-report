# -*- coding: utf-8 -*-
"""인스타그램(Meta) 광고 원본(Raw Data Report) 파서.

Meta Ads Manager 내보내기는 'Formatted Report'/'Raw Data Report' 두 시트를 갖거나
'Raw Data Report' 한 시트만 가질 수 있다. 광고 세트 단위로 뭉쳐진 'Formatted
Report'는 캠페인명이 비어 있는 등 콘텐츠 단위 매칭에 못 쓰므로 항상
'Raw Data Report'를 선택한다(LX세미콘 규칙과 동일).

내보내기마다 컬럼 순서가 달라(예: '도달'/'노출' 순서가 파일마다 다름) 위치가 아니라
헤더 이름으로 값을 찾는다. 한 행 = 한 광고(=대개 콘텐츠 1건)이며, 캠페인 이름의
접미사(_참여/_도달/_동영상 조회)로 광고 유형을 분류한다.
"""
from dataclasses import dataclass
import datetime as dt
import re
import openpyxl

_AD_NAME_DATE_RE = re.compile(r"^(\d{2})(\d{2})(\d{2})_")


def _date_from_ad_name(ad_name: str):
    """광고 이름 앞의 'YYMMDD_' 코드에서 발행일을 뽑는다.

    Meta 내보내기 형식에 따라 개별 광고 단위 '시작' 컬럼이 없는 경우가 있어(전체
    기간만 나옴), 광고 이름에 박혀 있는 날짜 코드를 발행일의 1차 근거로 쓴다.
    """
    m = _AD_NAME_DATE_RE.match(ad_name or "")
    if not m:
        return None
    yy, mm, dd = (int(x) for x in m.groups())
    try:
        return dt.date(2000 + yy, mm, dd)
    except ValueError:
        return None


@dataclass
class AdRow:
    campaign_name: str
    ad_name: str
    ad_type: str          # '참여' | '도달' | '동영상조회' | None(분류 불가)
    anchor_date: dt.date       # 광고 소재명의 날짜 코드 - 콘텐츠 매칭용 근사값(실제 발행일이 아닐 수 있음)
    campaign_start: dt.date    # 원본에 개별 광고 단위 '시작' 컬럼이 있을 때만 채워짐
    campaign_end: dt.date      # 위와 동일, '종료' 컬럼
    result_type: str
    result: float
    cost_per_result: float
    reach: float
    impressions: float
    video_2s_views: float
    video_2s_cost: float
    link_clicks: float
    cpc: float
    ctr: float
    cpm: float
    spend: float


def _classify(campaign_name: str) -> str:
    name = campaign_name or ""
    if "동영상" in name:
        return "동영상조회"
    if "참여" in name:
        return "참여"
    if "도달" in name:
        return "도달"
    return None


def _to_date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if v in (None, ""):
        return None
    try:
        return dt.datetime.strptime(str(v).strip(), "%Y-%m-%d").date()
    except ValueError:
        return None


def _to_num(v):
    if v in (None, "", "-"):
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def parse_ad_report(file_like) -> list:
    wb = openpyxl.load_workbook(file_like, data_only=True)
    sheet_name = "Raw Data Report" if "Raw Data Report" in wb.sheetnames else wb.sheetnames[0]
    ws = wb[sheet_name]

    header_row = None
    header = {}
    for r in range(1, min(ws.max_row, 10) + 1):
        values = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
        if "캠페인 이름" in values:
            header_row = r
            header = {v: c for c, v in enumerate(values, start=1) if v}
            break
    if header_row is None:
        raise ValueError("광고 원본에서 '캠페인 이름' 헤더를 찾을 수 없음")

    def get(row, *names):
        for name in names:
            c = header.get(name)
            if c:
                return ws.cell(row=row, column=c).value
        return None

    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        campaign = get(r, "캠페인 이름")
        if not campaign:
            continue
        ad_name = str(get(r, "광고 이름") or "").strip()
        if ad_name == "All":
            # 캠페인 전체를 뭉친 집계 행 - 개별 광고(콘텐츠) 행과 중복되므로 제외
            continue
        explicit_start = _to_date(get(r, "시작"))
        explicit_end = _to_date(get(r, "종료"))
        anchor = (_date_from_ad_name(ad_name) or explicit_start
                  or _to_date(get(r, "보고 시작")))
        rows.append(AdRow(
            campaign_name=str(campaign).strip(),
            ad_name=ad_name,
            ad_type=_classify(campaign),
            anchor_date=anchor,
            campaign_start=explicit_start,
            campaign_end=explicit_end,
            result_type=get(r, "결과 유형"),
            result=_to_num(get(r, "결과")),
            cost_per_result=_to_num(get(r, "결과당 비용")),
            reach=_to_num(get(r, "도달")),
            impressions=_to_num(get(r, "노출")),
            video_2s_views=_to_num(get(r, "동영상 연속 2초 이상 재생")),
            video_2s_cost=_to_num(get(r, "동영상 연속 2초 이상 재생당 비용")),
            link_clicks=_to_num(get(r, "링크 클릭")),
            cpc=_to_num(get(r, "CPC(링크 클릭당 비용)", "CPC(전체)")),
            ctr=_to_num(get(r, "CTR(링크 클릭률)", "CTR(전체)")),
            cpm=_to_num(get(r, "CPM(1,000회 노출당 비용)")),
            spend=_to_num(get(r, "지출 금액 (KRW)")),
        ))
    return rows
