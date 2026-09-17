# -*- coding: utf-8 -*-
"""전월 로우데이터의 시트 구조를 유지하면서 이번 달 데이터 블록을 이어붙이는 엔진.

각 콘텐츠/AD 시트는 공통적으로 다음 구조를 가진다.
  - 상단: 월별(7~12월) 요약 표
  - 중간: 데이터 테이블 헤더 행 ('NO.' 라벨)
  - 하단: 월별 데이터 블록(N행) + 평균 행 + 합계 행이 월 순서대로 반복

새 블록은 마지막 '합계' 행 바로 다음에 삽입하고, 스타일은 바로 위 데이터 블록의
마지막 데이터 행에서 복사한다. '추측 금지' 원칙에 따라 값이 없으면 빈 칸이 아니라
'확인 필요'를 직접 쓴다.
"""
import copy
from openpyxl.utils import get_column_letter
from openpyxl.formula.translate import Translator

CONFIRM_NEEDED = "확인 필요"


def find_header_row(ws, label, col, search_rows=30):
    for r in range(1, min(ws.max_row, search_rows) + 1):
        if ws.cell(row=r, column=col).value == label:
            return r
    return None


def find_last_label_row(ws, label, col, search_from=1, search_to=None):
    """지정한 열에서 label과 값이 같은 마지막 행 번호를 찾는다(예: 마지막 '합계' 행)."""
    search_to = search_to or ws.max_row
    last = None
    for r in range(search_from, search_to + 1):
        if ws.cell(row=r, column=col).value == label:
            last = r
    return last


def copy_row_style(ws, src_row, dst_row, min_col, max_col):
    for c in range(min_col, max_col + 1):
        src = ws.cell(row=src_row, column=c)
        dst = ws.cell(row=dst_row, column=c)
        if src.has_style:
            dst.font = copy.copy(src.font)
            dst.border = copy.copy(src.border)
            dst.fill = copy.copy(src.fill)
            dst.number_format = src.number_format
            dst.alignment = copy.copy(src.alignment)
            dst.protection = copy.copy(src.protection)


def insert_block(ws, insert_at, n_rows, style_row, min_col, max_col):
    """insert_at 위치에 n_rows개의 빈 행을 삽입하고 style_row 서식을 복제한다."""
    ws.insert_rows(insert_at, n_rows)
    # style_row가 삽입 지점보다 아래였다면 삽입만큼 밀려났으므로 보정
    src = style_row if style_row < insert_at else style_row + n_rows
    for i in range(n_rows):
        copy_row_style(ws, src, insert_at + i, min_col, max_col)


def write_row(ws, row, col_start, values):
    for i, v in enumerate(values):
        ws.cell(row=row, column=col_start + i, value=v)


def col_letter(idx):
    return get_column_letter(idx)


def range_formula(func, col_idx, r1, r2):
    """예: range_formula('SUM', 8, 23, 33) -> '=SUM(H23:H33)'."""
    L = get_column_letter(col_idx)
    return f"={func}({L}{r1}:{L}{r2})"


def copy_month_row_formula(ws, template_row, target_row, min_col, max_col):
    """월별 요약 표(예: 인스타그램 시트 6~11행)처럼, 같은 서식의 수식이 아래로
    드래그된 것과 동일한 패턴일 때(고정 범위 $C$14:$C$100 등은 그대로, 상대참조
    C6→C7만 이동) 템플릿 행의 수식을 target_row 기준으로 번역해서 옮긴다.
    수식이 아닌 셀(값/빈칸)은 건드리지 않는다."""
    for c in range(min_col, max_col + 1):
        src = ws.cell(row=template_row, column=c)
        if isinstance(src.value, str) and src.value.startswith("="):
            translated = Translator(src.value, origin=src.coordinate).translate_formula(
                ws.cell(row=target_row, column=c).coordinate)
            ws.cell(row=target_row, column=c, value=translated)
